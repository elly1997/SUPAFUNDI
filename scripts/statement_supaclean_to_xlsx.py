"""Convert Statement Supaclean 2025.pdf (NMB bank statement) to Excel."""
from __future__ import annotations

import json
import re
import shutil
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import easyocr
import openpyxl
import pdfplumber
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PDF_PATH = Path(r"c:\Users\HP\OneDrive\Desktop\Statement Supaclean 2025.pdf")
OUT_PATH = Path(r"c:\Users\HP\OneDrive\Desktop\Statement Supaclean 2025.xlsx")
DPI = 200
PAGE_WIDTH = 1653  # reference width at 200 DPI

# Column x-ranges as fractions of page width (calibrated from NMB statement)
COL_FRAC = [
    ("Book Date", 0.0, 0.088),
    ("Value Date", 0.088, 0.142),
    ("Trn Br Name", 0.142, 0.224),
    ("Narration", 0.224, 0.411),
    ("Xref", 0.411, 0.574),
    ("Cheque No", 0.574, 0.638),
    ("Debit", 0.638, 0.711),
    ("Credit", 0.711, 0.786),
    ("Balance", 0.786, 1.0),
]

TXN_CODE_RE = re.compile(r"\b(006|121|371|372|455|605)\b")

HEADERS = [c[0] for c in COL_FRAC]

DATE_RE = re.compile(r"(\d{2})[/.](\d{2})[/.](\d{4})")
AMT_RE = re.compile(r"^[\d,]+\.?\d*$")
BALANCE_RE = re.compile(r"^[\d,]+\.\d{2}$")
SKIP_PATTERNS = (
    "printed date",
    "printed by",
    "page number",
    "customer account statement",
    "book date",
    "value date",
    "narration",
    "opening balance",
    "total credit",
    "number of debit",
    "number of credit",
    "current balance",
    "uncollected",
    "available balance",
    "nmb bank plc",
    "received",
)


@dataclass
class Transaction:
    book_date: str = ""
    value_date: str = ""
    branch: str = ""
    narration: str = ""
    xref: str = ""
    cheque_no: str = ""
    debit: float | None = None
    credit: float | None = None
    balance: float | None = None
    page: int = 0
    notes: str = ""


@dataclass
class PageMeta:
    page_num: int
    title: str = "NMB CUSTOMER ACCOUNT STATEMENT — SUPACLEAN (2025)"
    summary: dict = field(default_factory=dict)


def normalize_date(text: str) -> str:
    m = DATE_RE.search(text.replace(" ", ""))
    if not m:
        return ""
    d, mo, y = m.groups()
    try:
        day, month, year = int(d), int(mo), int(y)
    except ValueError:
        return ""
    if year != 2025 or month < 1 or month > 12 or day < 1 or day > 31:
        return ""
    return f"{d.zfill(2)}/{mo.zfill(2)}/{y}"


def parse_amount(text: str) -> float | None:
    cleaned = text.strip().replace(" ", "")
    for old, new in (("O", "0"), ("o", "0"), ("d", "0"), ("D", "0"), ("l", "1"), ("I", "1"), ("S", "5")):
        cleaned = cleaned.replace(old, new)
    cleaned = cleaned.replace(",.", ",").replace("..", ".")
    if not cleaned or cleaned in {"0", "0.00", "-"}:
        return None
    if not re.match(r"^[\d,]+\.?\d*$", cleaned):
        return None
    try:
        val = float(cleaned.replace(",", ""))
    except ValueError:
        return None
    return val if val != 0 else None


def assign_column(x: float, page_width: float) -> str:
    frac = x / page_width if page_width else x / PAGE_WIDTH
    for name, x0, x1 in COL_FRAC:
        if x0 <= frac < x1:
            return name
    return "Narration"


def cluster_lines(
    items: list[tuple[float, float, str, float]], page_width: float, y_tol: float = 14
) -> list[dict]:
    """Group OCR boxes into horizontal lines with column assignments."""
    if not items:
        return []
    sorted_items = sorted(items, key=lambda t: (t[1], t[0]))
    lines: list[dict] = []
    current_y = sorted_items[0][1]
    current: dict[str, list[str]] = {h: [] for h in HEADERS}
    current_y_vals: list[float] = []

    def flush() -> None:
        nonlocal current, current_y_vals
        if any(current[c] for c in HEADERS):
            lines.append({k: " ".join(v).strip() for k, v in current.items()})
        current = {h: [] for h in HEADERS}
        current_y_vals = []

    for x, y, text, _conf in sorted_items:
        if abs(y - current_y) > y_tol:
            flush()
            current_y = y
        col = assign_column(x, page_width)
        current[col].append(text)
        current_y_vals.append(y)

    flush()
    return lines


def is_noise_line(line: dict) -> bool:
    joined = " ".join(line.values()).lower()
    if any(p in joined for p in SKIP_PATTERNS):
        return True
    if joined.strip() in {"nmb", "0", ""}:
        return True
    if "account number" in joined or "customer no" in joined or "supacleanl" in joined:
        return True
    return False


def is_valid_transaction(txn: Transaction) -> bool:
    if txn.balance is None:
        return False
    if not txn.narration and not txn.book_date:
        return False
    if txn.narration and not TXN_CODE_RE.search(txn.narration) and not txn.book_date:
        return False
    header_bits = ("account", "customcr", "branch id", "currency", "001683445", "4086600442")
    narr = txn.narration.lower()
    if any(b in narr for b in header_bits):
        return False
    return True


def merge_narration(parts: list[str]) -> str:
    return " ".join(p for p in parts if p).strip()


def lines_to_transactions(lines: list[dict], page_num: int) -> list[Transaction]:
    """Parse clustered lines into transaction records."""
    txns: list[Transaction] = []
    buffer: list[dict] = []

    def finalize_block(block: list[dict]) -> Transaction | None:
        if not block:
            return None
        txn = Transaction(page=page_num)
        narr_parts: list[str] = []

        for ln in block:
            for col in HEADERS:
                val = ln.get(col, "").strip()
                if not val:
                    continue
                d = normalize_date(val)
                if d:
                    if not txn.book_date:
                        txn.book_date = d
                    elif not txn.value_date and d != txn.book_date:
                        txn.value_date = d
                    elif not txn.value_date:
                        txn.value_date = d
            if ln.get("Trn Br Name"):
                br = ln["Trn Br Name"].strip()
                if br and br not in {"0", "Office", "Olfico", "Clfice", "Hcad"}:
                    if not txn.branch:
                        txn.branch = br
                    elif br not in txn.branch:
                        txn.branch = f"{txn.branch} {br}".strip()
            for key in ("Narration", "Xref", "Cheque No"):
                val = ln.get(key, "").strip()
                if val and val not in {"0"}:
                    if key == "Narration":
                        narr_parts.append(val)
                    elif key == "Xref" and not txn.xref:
                        txn.xref = val
                    elif key == "Cheque No" and not txn.cheque_no:
                        txn.cheque_no = val.replace(" ", "")
            debit = parse_amount(ln.get("Debit", ""))
            credit = parse_amount(ln.get("Credit", ""))
            balance_raw = ln.get("Balance", "").strip()
            balance = parse_amount(balance_raw) if BALANCE_RE.match(balance_raw.replace(" ", "")) else None
            if debit is not None:
                txn.debit = debit
            if credit is not None:
                txn.credit = credit
            if balance is not None:
                txn.balance = balance

        txn.narration = merge_narration(narr_parts)
        # Fallback: scan narration/xref for dates missed by column assignment
        if not txn.book_date:
            for part in narr_parts + [txn.xref]:
                d = normalize_date(part)
                if d:
                    txn.book_date = d
                    break
        if not txn.book_date and not txn.balance and not txn.debit and not txn.credit:
            return None
        if txn.book_date and not txn.value_date:
            txn.value_date = txn.book_date
        return txn

    for ln in lines:
        if is_noise_line(ln):
            continue
        has_balance = bool(ln.get("Balance") and BALANCE_RE.match(ln["Balance"].replace(" ", "")))
        has_amounts = parse_amount(ln.get("Debit", "")) or parse_amount(ln.get("Credit", ""))
        starts_txn = bool(normalize_date(ln.get("Book Date", "")) or normalize_date(ln.get("Value Date", "")))

        if has_balance or (has_amounts and starts_txn):
            buffer.append(ln)
            txn = finalize_block(buffer)
            if txn and is_valid_transaction(txn):
                txns.append(txn)
            buffer = []
        else:
            buffer.append(ln)

    if buffer:
        txn = finalize_block(buffer)
        if txn and is_valid_transaction(txn):
            txns.append(txn)

    return txns


def ocr_page(
    reader: easyocr.Reader, image_path: Path, page_num: int
) -> tuple[list[tuple[float, float, str, float]], float]:
    from PIL import Image

    with Image.open(image_path) as im:
        page_width = im.width
    results = reader.readtext(str(image_path), detail=1, paragraph=False)
    y_min = int(page_width * 0.21) if page_num == 1 else int(page_width * 0.145)
    items: list[tuple[float, float, str, float]] = []
    for bbox, text, conf in results:
        if conf < 0.2 or not text.strip():
            continue
        y = (bbox[0][1] + bbox[2][1]) / 2
        x = (bbox[0][0] + bbox[1][0]) / 2
        if y < y_min:
            continue
        items.append((x, y, text.strip(), conf))
    return items, float(page_width)


def extract_summary(page_items: list[tuple[float, float, str, float]]) -> dict:
    text = " ".join(t[2] for t in page_items).lower()
    summary: dict = {}

    def find_amount(label: str) -> float | None:
        idx = text.find(label)
        if idx == -1:
            return None
        snippet = text[idx : idx + 80]
        nums = re.findall(r"[\d,]+\.?\d*", snippet)
        for n in nums:
            val = parse_amount(n)
            if val is not None:
                return val
        return None

    for label, key in [
        ("total credit amount", "total_credit"),
        ("current balance", "current_balance"),
        ("available balance", "available_balance"),
        ("uncollected amount", "uncollected"),
    ]:
        val = find_amount(label)
        if val is not None:
            summary[key] = val

    for label, key in [
        ("number of debit transactions", "debit_count"),
        ("number of credit transactions", "credit_count"),
    ]:
        idx = text.find(label)
        if idx != -1:
            nums = re.findall(r"\d+", text[idx : idx + 60])
            if nums:
                summary[key] = int(nums[0])

    return summary


# Excel styling (matches stock list scripts)
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor="1E293B")
title_font = Font(bold=True, size=14, color="F97316")
debit_fill = PatternFill("solid", fgColor="FEE2E2")
credit_fill = PatternFill("solid", fgColor="D1FAE5")
thin = Side(style="thin", color="334155")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)
right = Alignment(horizontal="right", vertical="center")


def txn_to_row(txn: Transaction) -> list:
    return [
        txn.book_date,
        txn.value_date,
        txn.branch,
        txn.narration,
        txn.xref,
        txn.cheque_no,
        txn.debit,
        txn.credit,
        txn.balance,
    ]


def write_sheet(ws, title: str, rows: list[list], start_row: int = 4) -> int:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws["A1"] = title
    ws["A1"].font = title_font
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 28

    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=start_row, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    for ri, row in enumerate(rows, start_row + 1):
        for ci, val in enumerate(row, 1):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.border = border
            if ci in (1, 2):
                cell.alignment = center
            elif ci in (7, 8, 9):
                cell.alignment = right
                if isinstance(val, (int, float)):
                    cell.number_format = "#,##0.00"
            elif ci == 4:
                cell.alignment = left
            else:
                cell.alignment = left
            if ci == 7 and val:
                cell.fill = debit_fill
            if ci == 8 and val:
                cell.fill = credit_fill

    total_row = start_row + len(rows) + 1
    ws.cell(row=total_row, column=6, value="TOTAL:").font = Font(bold=True)
    ws.cell(row=total_row, column=6).alignment = Alignment(horizontal="right")
    for col, letter in [(7, "G"), (8, "H")]:
        c = ws.cell(
            row=total_row,
            column=col,
            value=f"=SUM({letter}{start_row + 1}:{letter}{total_row - 1})",
        )
        c.font = Font(bold=True)
        c.number_format = "#,##0.00"
        c.alignment = right

    widths = [12, 12, 18, 55, 18, 12, 14, 14, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    return total_row


def build_workbook(
    all_txns: list[Transaction],
    pages: dict[int, list[Transaction]],
    summary: dict,
    page_count: int,
) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    doc_title = "NMB CUSTOMER ACCOUNT STATEMENT — SUPACLEAN (Jan–Dec 2025)"

    all_rows = [txn_to_row(t) for t in all_txns]
    ws_all = wb.create_sheet("All Records", 0)
    write_sheet(ws_all, f"{doc_title} (All Pages)", all_rows)

    for page_num in sorted(pages):
        if page_num == 74:
            continue
        ws = wb.create_sheet(title=f"Page {page_num}")
        write_sheet(ws, f"{doc_title} — Page {page_num}", [txn_to_row(t) for t in pages[page_num]])

    ws_sum = wb.create_sheet("Summary")
    ws_sum["A1"] = "Summary"
    ws_sum["A1"].font = title_font

    total_debit = sum(t.debit or 0 for t in all_txns)
    total_credit = sum(t.credit or 0 for t in all_txns)
    debit_txns = sum(1 for t in all_txns if t.debit)
    credit_txns = sum(1 for t in all_txns if t.credit)
    closing = all_txns[-1].balance if all_txns else None

    summary_data = [
        ["Document Type", "Bank Statement (NMB)"],
        ["Account", "SUPACLEAN — 4086600442"],
        ["Statement Period", "01/01/2025 – 31/12/2025"],
        ["Total Pages", page_count],
        ["Total Transactions", len(all_txns)],
        ["Total Debit (TZS)", total_debit],
        ["Total Credit (TZS)", total_credit],
        ["Debit Transactions", debit_txns],
        ["Credit Transactions", credit_txns],
        ["Closing Balance (TZS)", closing],
    ]
    if summary.get("total_credit"):
        summary_data.append(["PDF Total Credit (TZS)", summary["total_credit"]])
    if summary.get("current_balance"):
        summary_data.append(["PDF Current Balance (TZS)", summary["current_balance"]])
    if summary.get("debit_count"):
        summary_data.append(["PDF Debit Count", summary["debit_count"]])
    if summary.get("credit_count"):
        summary_data.append(["PDF Credit Count", summary["credit_count"]])

    for i, (label, val) in enumerate(summary_data, 3):
        ws_sum.cell(row=i, column=1, value=label).font = Font(bold=True)
        c = ws_sum.cell(row=i, column=2, value=val)
        if isinstance(val, float) or (isinstance(val, int) and i > 5):
            c.number_format = "#,##0.00" if isinstance(val, float) or val > 999 else "0"
    ws_sum.column_dimensions["A"].width = 28
    ws_sum.column_dimensions["B"].width = 22

    return wb


def save_cache(
    cache_path: Path,
    all_txns: list[Transaction],
    pages: dict[int, list[Transaction]],
    pdf_summary: dict,
    page_count: int,
    last_done: int,
) -> None:
    cache_path.write_text(
        json.dumps(
            {
                "page_count": page_count,
                "last_done": last_done,
                "summary": pdf_summary,
                "pages": {str(k): len(v) for k, v in pages.items()},
                "transactions": [
                    {
                        "page": t.page,
                        "book_date": t.book_date,
                        "value_date": t.value_date,
                        "branch": t.branch,
                        "narration": t.narration,
                        "xref": t.xref,
                        "cheque_no": t.cheque_no,
                        "debit": t.debit,
                        "credit": t.credit,
                        "balance": t.balance,
                    }
                    for t in all_txns
                ],
            },
            indent=2,
        ),
        encoding="utf-8",
    )


def load_cache(cache_path: Path) -> tuple[list[Transaction], dict[int, list[Transaction]], dict, int, int] | None:
    if not cache_path.exists():
        return None
    try:
        data = json.loads(cache_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return None
    txns = [
        Transaction(
            page=d["page"],
            book_date=d.get("book_date", ""),
            value_date=d.get("value_date", ""),
            branch=d.get("branch", ""),
            narration=d.get("narration", ""),
            xref=d.get("xref", ""),
            cheque_no=d.get("cheque_no", ""),
            debit=d.get("debit"),
            credit=d.get("credit"),
            balance=d.get("balance"),
        )
        for d in data.get("transactions", [])
    ]
    pages: dict[int, list[Transaction]] = {}
    for t in txns:
        pages.setdefault(t.page, []).append(t)
    for page_num, count in data.get("pages", {}).items():
        pn = int(page_num)
        if pn not in pages:
            pages[pn] = []
        elif count == 0 and pn not in {t.page for t in txns}:
            pages[pn] = []
    return txns, pages, data.get("summary", {}), data.get("page_count", 0), data.get("last_done", 0)


def main() -> None:
    if not PDF_PATH.exists():
        raise FileNotFoundError(f"PDF not found: {PDF_PATH}")

    tmp_dir = Path(tempfile.gettempdir()) / "supaclean_stmt_active"
    tmp_dir.mkdir(exist_ok=True)
    cache_path = tmp_dir / "ocr_cache.json"
    img_dir = tmp_dir / "pages"
    img_dir.mkdir(exist_ok=True)

    print(f"Temp dir: {tmp_dir}", flush=True)
    print("Loading EasyOCR...", flush=True)
    reader = easyocr.Reader(["en"], gpu=False, verbose=False)

    all_txns: list[Transaction] = []
    pages: dict[int, list[Transaction]] = {}
    pdf_summary: dict = {}
    resume_from = 0

    cached = load_cache(cache_path)
    if cached:
        all_txns, pages, pdf_summary, cached_pages, resume_from = cached
        print(f"Resuming from page {resume_from + 1} ({len(all_txns)} txns cached)", flush=True)

    with pdfplumber.open(PDF_PATH) as pdf:
        page_count = len(pdf.pages)
        print(f"Pages: {page_count}", flush=True)

        # Quick text probe — confirm image-only PDF
        sample_text = (pdf.pages[0].extract_text() or "") + (pdf.pages[1].extract_text() or "")
        if len(sample_text.strip()) > 50:
            print("Text layer detected; using pdfplumber (not expected for this PDF)", flush=True)
        else:
            print("Image-only PDF — using OCR extraction", flush=True)

        for idx, page in enumerate(pdf.pages):
            page_num = idx + 1
            if page_num <= resume_from:
                continue

            img_path = img_dir / f"page_{page_num:03d}.png"
            if not img_path.exists():
                page.to_image(resolution=DPI).save(img_path, format="PNG")

            items, page_width = ocr_page(reader, img_path, page_num)
            lines = cluster_lines(items, page_width)
            if page_num == page_count:
                pdf_summary = extract_summary(items)
                pages[page_num] = []
                print(f"  Page {page_num}: summary page", flush=True)
            else:
                txns = lines_to_transactions(lines, page_num)
                pages[page_num] = txns
                all_txns.extend(txns)
                print(f"  Page {page_num}: {len(txns)} transactions", flush=True)

            save_cache(cache_path, all_txns, pages, pdf_summary, page_count, page_num)

    wb = build_workbook(all_txns, pages, pdf_summary, page_count)
    wb.save(OUT_PATH)

    total_debit = sum(t.debit or 0 for t in all_txns)
    total_credit = sum(t.credit or 0 for t in all_txns)
    ambiguous = [t for t in all_txns if not t.book_date or (not t.debit and not t.credit)]

    print(f"\nSaved: {OUT_PATH}", flush=True)
    print(f"Document type: Bank Statement (NMB Customer Account Statement)", flush=True)
    print(f"Pages: {page_count}", flush=True)
    print(f"Total transactions: {len(all_txns)}", flush=True)
    print(f"Total debit: TZS {total_debit:,.2f}", flush=True)
    print(f"Total credit: TZS {total_credit:,.2f}", flush=True)
    if all_txns:
        print(f"Closing balance: TZS {all_txns[-1].balance:,.2f}" if all_txns[-1].balance else "", flush=True)
    if pdf_summary:
        print(f"PDF summary: {pdf_summary}", flush=True)
    print(f"Ambiguous/incomplete rows: {len(ambiguous)}", flush=True)
    if ambiguous[:5]:
        for t in ambiguous[:5]:
            print(f"  Page {t.page}: {t.narration[:80]}...", flush=True)

    shutil.rmtree(tmp_dir, ignore_errors=True)
    print("Temp files cleaned up.", flush=True)


if __name__ == "__main__":
    main()
