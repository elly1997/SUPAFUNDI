"""Build Supafundi import spreadsheets from scanned PDF inventory lists."""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(r"c:\Users\HP\OneDrive\Desktop\SUPAFUNDI_Inventory_From_Scans.xlsx")

HEADERS = [
    "Page",
    "Section",
    "Code",
    "Brand",
    "Item",
    "Category / Spec",
    "Quantity",
    "Cost / Buying (TZS)",
    "Retail / Selling (TZS)",
    "Unit",
    "Notes",
]


def row(
    page: str,
    section: str,
    item: str,
    *,
    code: str = "",
    brand: str = "",
    category: str = "",
    qty: str = "",
    cost: str = "",
    retail: str = "",
    unit: str = "",
    notes: str = "",
) -> list:
    return [page, section, code, brand, item, category, qty, cost, retail, unit, notes]


PLUMBING: list[list] = [
    # Page 1 — PLUMBING SHELF 10/05/2026
    row("1", "PLUMBING SHELF (10/05/2026)", "Sample Product", code="PRD001", category="General", qty="10", cost="1000", unit="pcs"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Corner Gutter L", qty="3", cost="4500", retail="8000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Corner Gutter R.", qty="12", cost="4500", retail="8000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "PVC - 110mm"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Elbow 45°", category="Plumbing", qty="26pcs", cost="4500", retail="8000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Outlet inch 2.5", cost="4500", retail="8000", notes="Retail crossed out on scan"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Elbow 90°", category="plumbing", qty="16"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Outlet 3\"", qty="5pcs", cost="4500", retail="8000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Outlet 2.5\"", qty="8pcs", cost="4500", retail="8000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Stand za gas", qty="4", retail="10000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Tank 2000L", qty="3", cost="300000", retail="350000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Pvc Elbow 4\"", qty="2", cost="2000", retail="3500"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Brush roller 9\"", qty="134", cost="1500", retail="2000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Brush 5\" TBP", brand="China", qty="26", cost="3000", retail="5000", notes="Cost crossed out"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Brush 4\" TBP", qty="10", cost="4000", retail="4500"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Brush 5\" TBP", qty="4", cost="5000", retail="5500"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Brush 2\" TBP", qty="22", cost="2000", retail="2500"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Brush 3\" TBP", qty="11", cost="3000", retail="3500"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Brush 1\" TBP", qty="11", cost="1000", retail="1500"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Brush china 2\"", qty="6", cost="1000", retail="2000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Brush china 2.5\"", qty="4", cost="1000", retail="2500"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Brush china 1 ½\"", qty="8", cost="800", retail="1500"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Brush 3\"", qty="1", cost="2000", retail="3000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Tembo Pana", qty="2", cost="21000", retail="25000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Stop End R.", qty="40", cost="2500", retail="2500"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Stop End L", qty="23", cost="2500", retail="2500"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Brush Roller 2\"", qty="73", cost="1000", retail="2000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Kamba kudy Ndogo", qty="40", cost="1000", retail="2000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Kamba Manilla", qty="44", cost="600", retail="1000", notes="Cost crossed out"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Down pipe clamp", qty="123", cost="1300", retail="2500"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Gutter Clamp", qty="59", cost="1300", retail="1500"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Connector Gutter", qty="5", cost="2500", retail="4500"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Shoka", qty="7+1", cost="18000", retail="25000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Panga Ndogo", qty="9+1", cost="3000", retail="6000", notes="Retail had 8 crossed out"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Kyamba", qty="10", cost="2540", retail="5000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Panga kubwa", qty="16+3", cost="3500", retail="7000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Pad Msasa 7\"", qty="19", cost="5000", retail="7000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Pad 4\"", qty="14", cost="2500", retail="5000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Deco Kibao Nyeusi", qty="7", cost="5000", retail="8000"),
    row("1", "PLUMBING SHELF (10/05/2026)", "Stopper Kubwa", qty="56", cost="1500", retail="3500"),
    # Page 2
    row("2", "General items", "Kamru", qty="20", cost="1500", retail="2500", unit="pcs"),
    row("2", "General items", "Kibno plasta", qty="8", cost="5000"),
    row("2", "General items", "Mwiko 22\"", qty="5", cost="4000", retail="6000"),
    row("2", "General items", "Mwiko 20\"", qty="6", cost="3000", retail="5000"),
    row("2", "General items", "Mwiko 6\"", qty="8", cost="1500", retail="3000"),
    row("2", "General items", "Mwiko 7'", qty="4", cost="2000", retail="4000"),
    row("2", "General items", "Karamala Flat", qty="11", cost="6500", retail="10000"),
    row("2", "General items", "Karamala Dekok", qty="4", cost="4500", retail="7000"),
    row("2", "General items", "Scrapper 8\" JECO", qty="12", cost="2000", retail="3000 / 4000", notes="Two selling prices noted"),
    row("2", "General items", "Scrapper 10\" JECO", qty="13", cost="3000", retail="5000"),
    row("2", "General items", "Shelf bracket 10\" x 12\"", qty="50"),
    row("2", "General items", "Shelf brcts 8\" x 10\"", qty="49"),
    row("2", "General items", "Shelf brcts 6\" x 8\"", qty="51"),
    row("2", "General items", "Socket Pipe 2.5\"", qty="2pcs"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "Non Return Valve 1\"", notes="Crossed out"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "Non Return valve 3/4\"", notes="Crossed out"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "PPR Tee Plain 1\"", qty="32", cost="1800", retail="2000", unit="pcs"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "IPS Tee Plain 1\"", qty="7", cost="1000", retail="2000", unit="pcs"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "PPR Female Elbow 3/4 * 1", qty="18", cost="3500", retail="5000", unit="pcs"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "PPR Male Elbow 3/4 * 1", qty="19", cost="3500", retail="5000", unit="pcs"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "PPR Female socket 1 * 1", qty="9", cost="3500", retail="5000", unit="pcs"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "PPR Female socket 1 * 3/4", qty="9", cost="3500", retail="5000", unit="pcs"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "PPR Female socket 3/4 * 1/2", qty="17", cost="2500", retail="3500", unit="pcs"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "PPR Female socket 3/4 * 3/4", qty="6", cost="2500", retail="3500", unit="pcs"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "PPR Male socket 1 * 1", qty="14", cost="3500", retail="5000", unit="pcs"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "PPR Male socket 1 * 3/4", qty="10", cost="3500", retail="5000", unit="pcs"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "PPR Male socket 3/4 * 1/2", qty="53", cost="2500", retail="3500", unit="pcs"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "IPS Plain Elbow 1\"", qty="42", cost="1000", retail="2000", unit="pcs"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "PPR Plain Elbow 1\"", qty="26", cost="1000", retail="2000", unit="pcs"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "PPR Plain Socket 3/4", qty="20", cost="2500", retail="3500", unit="pcs"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "IPS Plain Socket 3/4", qty="83", cost="500", retail="1500", unit="per"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "GS Socket 3/4", qty="24", cost="1000", retail="2000", unit="pcs"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "IPS Plain Elbow 3/4", qty="91", cost="500", retail="1500", unit="pcs"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "GS Elbow 3/4", qty="17", cost="1000", retail="2000", unit="pcs"),
    row("2", "Plumbing Items Front desk (12/08/2024)", "GS Elbow 1/2", qty="4", cost="800", retail="1500", unit="per"),
]

# Page 3 plumbing — continuation in script part 2
PLUMBING_P3 = [
    ("PPR Plain socket 1\"", "28", "1000", "2000"),
    ("IPS Plain socket 1\"", "48", "1000", "2000"),
    ("IPS R. Bush 1 1/4 * 1", "4", "3000", "5000"),
    ("IPS R. Bush 1 1/2 * 1", "20", "3000", "5000"),
    ("IPS R. Bush 1 1/2 * 1 1/4", "5", "3000", "5000"),
    ("IPS R. Bush 3/4 * 1/2", "1", "1000", "2000"),
    ("PPR Female socket 1/2 * 1/2", "9", "2500", "3500"),
    ("PPR Female socket 3/4 * 3/4", "27", "2500", "3500"),
    ("IPS Plug 1\"", "45", "1000", "2000"),
    ("IPS Plug 3/4\"", "19", "800", "1500"),
    ("IPS Plug 1/2\"", "64", "500", "1000"),
    ("PPR Plug 1/2\"", "45", "500", "1000"),
    ("GS Plug 3/4", "4", "1000", "2000"),
    ("IPS R. Socket 1/2 * 3/4", "52", "800", "1500"),
    ("IPS R. Socket 1 * 3/4", "13", "1000", "2000"),
    ("GS R. Socket 3/4 * 1/2", "18", "1000", "2000"),
    ("GS R. Socket 1 1/4 * 3/4", "1", "3500", "5000"),
    ("PPR Union 1\"", "10", "3000", "5000"),
    ("PPR Union 3/4", "1", "2500", "3500"),
    ("IPS Union 1\"", "2", "3000", "4500"),
    ("GS Union 3/4", "4", "2500", "4000"),
    ("GS R. Bush 1/2 * 3/4", "6", "1000", "2000"),
    ("IPS R. Bush 1/2 * 3/4", "37", "800", "2000"),
    ("IPS R. Bush 3/4 * 1", "19", "1000", "2500"),
    ("IPS Nipple 1\"", "19", "1000", "2000"),
    ("IPS Nipple socket 1/2", "71", "500", "1000"),
    ("GS Nipple socket 1/2", "63", "800", "1500"),
    ("GS Socket 1/2", "24", "800", "1500"),
    ("IPS Socket 1/2 * 1/2", "39", "500", "1000"),
    ("PPR Plain socket 1 * 3/4", "3", "1000", "2000"),
    ("PPR Plain socket 1/2 * 1/2", "47", "500", "1000"),
    ("PPR Plain socket 3/4 * 1/2", "37", "800", "1500"),
    ("GS Nipple 3/4", "21", "1000", "2000"),
    ("GS Nipple 1/2", "47", "800", "1500"),
    ("IPS Nipple 1/2", "112", "500", "1000"),
    ("IPS Nipple 3/4", "203", "800", "1500"),
    ("PPR Plain Elbow 1/2", "110", "500", "1000"),
    ("IPS Plain Elbow 1/2", "52", "500", "1000"),
    ("PPR Plain Tee 1/2", "54", "600", "1000"),
    ("IPS Plain Tee 1/2", "49", "600", "1000"),
    ("GS Tee 1/2", "15", "800", "1500"),
]
for item, qty, bp, sp in PLUMBING_P3:
    PLUMBING.append(row("3", "Plumbing fittings (cont.)", item, qty=qty, cost=bp, retail=sp, unit="pcs"))

PLUMBING_P4 = [
    ("", "PPR Tee Female 3/4 * 1/2", "45", "2500", "3500"),
    ("", "PPR Tee Male 3/4 * 1/2", "18", "2500", "3500"),
    ("", "PPR Tee Female 3/4 * 3/4", "6", "2500", "3500"),
    ("", "PPR Tee Female 1/2 * 1/2", "11", "2000", "3000"),
    ("", "IPS Tee 3/4", "116", "800", "1500"),
    ("", "GI Tee 3/4", "10", "1000", "2000"),
    ("", "PPR Plain Tee 3/4", "7", "800", "1500"),
    ("", "PPR Female Elbow 3/4 * 1/2", "40", "2500", "3500"),
    ("", "PPR Male Elbow 3/4 * 1/2", "15", "2500", "3500"),
    ("", "PPR Female Elbow 1/2 * 1/2", "13", "2000", "3000"),
    ("", "PPR Female Elbow 3/4 * 3/4", "47", "2500", "3500"),
    ("OG", "Bib cock PEX 3/4", "9", "15000", "20000"),
    ("OG", "Bib cock PEX 1/2", "4", "13000", "17000"),
    ("OG", "Gate valve PEX 3/4", "5", "15000", "20000"),
    ("", "PVC Tee 4\"", "21", "2000", "3500"),
    ("", "PVC Elbow 4\" 45°", "39", "2000", "3500"),
    ("", "PVC Elbow 4\" 90°", "38", "2000", "3500"),
    ("", "PVC (Inspection) band", "10", "3000", "5000"),
    ("", "Flush tank pipe bend", "15", "5000", "7000"),
    ("", "PVC Y Tee 4\"", "7", "5000", "8000"),
    ("", "Strainer - plastic", "41", "2500", "5000"),
    ("", "Vent cup 4\"", "12", "2000", "3000"),
    ("", "PVC Plug 4\"", "22", "2000", "3500"),
    ("", "R. Bush", "", "", "", "No qty/price on scan"),
    ("", "PVC Tee 1 1/2", "10", "1000", "2000"),
    ("", "PVC Elbow 1 1/2 90°", "66", "1000", "2000"),
    ("", "PVC Elbow 1 1/2 45°", "24", "1000", "2000"),
    ("", "PVC Elbow 2\" 90°", "27", "1200", "2500"),
    ("", "PVC Elbow 2\" 45°", "8", "1200", "2500"),
    ("", "PVC Tee 1 1/2", "26", "900", "1500"),
    ("", "PVC Elbow 1 1/4 90°", "15", "900", "1500"),
    ("", "PVC Elbow 1 1/4 45°", "13", "900", "1500"),
    ("", "PVC Tee 1 1/4", "49", "900", "1500"),
    ("", "PVC Bush 2\"", "25", "2000", "3000"),
    ("", "PVC Socket 1 1/2", "12", "1000", "2000"),
    ("", "PVC vent ndogo 2\"", "3", "1000", "2000"),
    ("", "PVC Plug 1 1/2", "16", "1000", "2000"),
    ("", "PVC Y 1 1/2", "11", "2000", "3500"),
    ("", "PVC plug 1 1/4", "22", "1000", "2000"),
]
for brand, item, qty, bp, sp, *extra in PLUMBING_P4:
    notes = extra[0] if extra else ("" if sp else "")
    if "No qty" in str(sp) or extra:
        notes = extra[0] if extra else "No qty/price on scan"
        sp = ""
    PLUMBING.append(
        row("4", "Plumbing & PVC (cont.)", item, brand=brand, qty=qty, cost=bp, retail=sp, unit="pcs", notes=notes)
    )

PLUMBING_P5 = [
    ("R. Bush 2\" * 2 1/2", "22", "2000", "3000"),
    ("R. Bush 2\" * 1 1/2", "22", "1000", "2000"),
    ("R. Bush 2\" * 1 1/4", "26", "1000", "2000"),
    ("R. Bush 1 1/2 * 1 1/4", "5", "1000", "2000"),
    ("R. Socket 1 1/2", "17", "", ""),
    ("Four way", "10", "8000", "10000"),
    ("Welding Bush 1/2\"", "450", "800/pair", "1000/pair", "@ 500"),
    ("Welding Bush 1\"", "250", "2400/pair", "3000/pair", "@ 1500"),
    ("Vitunguu Medium size", "73", "800", "1500"),
    ("Vitunguu Vidogo", "49", "800", "1200"),
    ("Welding Bush Bawa kubwa", "364", "3000/pair", "4000/pair", "@ 2000"),
    ("Welding Bush 3/4", "240", "1500/pair", "2000/pair", "@ 1000"),
    ("Welding Bush Bawa ndogo", "180", "2300/pair", "3000/pair", "@ 1500"),
    ("Metal Clamp 3/4", "318", "200", "500"),
    ("Metal Clamp 1\"", "210", "500", "1000"),
    ("Metal clamp 4\"", "26", "1000", "2000"),
    ("Plastic clamp 4\"", "7", "1000", "2000", "Mat crossed out"),
    ("PPR pipe 1/2", "18", "6000", "10000"),
    ("PPR pipe 3/4", "16", "8000", "15000"),
    ("PPR pipe 1\"", "15", "18000", "25000"),
    ("PPR pipe 1 1/4", "4", "25000", "35000"),
    ("IPS pipe 1/2", "6", "6000", "10000"),
    ("IPS pipe 3/4", "1", "8000", "15000"),
    ("IPS pipe 1\"", "1", "18000", "25000"),
    ("PVC Gutter ERA", "19", "23500", "28000"),
    ("PVC Gutter Metro", "13", "20000", "26000"),
    ("Nondo (Y16) mm", "26", "42400", "45000"),
    ("Nondo (Y12) mm", "74", "23800", "26000"),
    ("Nondo (Y8) mm", "23", "12500", "14000"),
    ("Nondo (Y10) mm", "45", "17600", "18500"),
    ("Wire mesh Tanzania", "25", "13500", "18000"),
    ("Wire mesh Kenya", "20", "7000", "12000"),
    ("Bati Kiboko 3m G30", "12", "17500", "23000"),
    ("Bati Kiboko 2.5m G30", "10", "14736", "20000"),
    ("Metal plate G22", "5", "39000", "45000"),
    ("Down pipe", "17", "20000", "26000", "BP 20000 crossed; 18000 noted"),
    ("Metal Angle line 1 1/2 * 1 1/2", "9", "30000", "33000"),
    ("Furniture pipe 1/2", "16", "7500", "10000"),
    ("Furniture pipe 3/4", "16", "9000", "12000"),
    ("Flat Bar 3/4", "14", "7000", "8500"),
]
for item, qty, bp, sp, *rest in PLUMBING_P5:
    PLUMBING.append(
        row("5", "Pipes, clamps & building", item, qty=qty, cost=bp, retail=sp, unit="pcs", notes=rest[0] if rest else "")
    )

PLUMBING_P6_TOP = [
    ("", "Flat Bar", "1\"", "10", "8000", "10000"),
    ("", "Flat Bar", "1 1/2 (3mm)", "11", "12500", "15000"),
    ("", "PVC pipe", "4\"", "44", "20000", "26000"),
    ("", "pvc pipe", "1 1/4", "24", "5000", "10000"),
    ("", "PVC pipe", "1 1/2", "42", "7000", "12000"),
    ("", "Hollow section", "1 1/2 * 1 1/2", "24", "14000", "18000"),
    ("", "Hollow section", "3/4 1.5mm", "34", "15000", "18000"),
    ("", "Hollow section", "1 1/4 1mm", "5", "11500", "14000"),
    ("", "Hollow section", "2\"", "18", "20000", "25000"),
    ("", "Hollow section", "1 1/2 * 1", "9", "12300", "16000"),
    ("", "Hollow section", "1 * 1", "20", "7800", "12000"),
    ("", "HDPE polly pipe", "3/4", "4", "100000", "150000", "roller"),
    ("", "HDPE polly pipe", "1\"", "1", "135000", "170000", "roller"),
]
for brand, item, spec, qty, bp, sp, *u in PLUMBING_P6_TOP:
    unit = u[0] if u else "pcs"
    PLUMBING.append(
        row("6", "Sections & pipes", item, brand=brand, category=spec, qty=qty, cost=bp, retail=sp, unit=unit)
    )

PLUMBING_P6_GEN = [
    ("", "Allan Key", "6mm", "25"),
    ("", "Alan key", "5mm", "23"),
    ("", "Alan Key", "4mm", "25"),
    ("", "Alan key", "3mm", "10"),
    ("", "Alan key", "3.5mm", "10"),
    ("", "Alan key", "2.5mm", "9"),
    ("", "Alan key", "1.5mm", "45"),
    ("Andika", "Wood Bit", "35mm", "5", "10000", "15000"),
    ("", "Wood screw", "1 1/4\"", "2", "2000", "3500", "boxes"),
    ("", "Wood screw", "1 1/2\"", "1", "2000", "3500", "box"),
    ("", "Bolt za vitasa", "2 1/2\"", "90", "350", "500"),
    ("", "Star bit", "(King Lion)", "17", "1500", "2500"),
    ("", "Star Bit", "(Jumlee)", "18", "1000", "2000"),
    ("", "Steel Bit", "10mm", "20", "7000", "10000"),
    ("", "Golden bolt", "2\"", "90", "200", "500"),
    ("", "Raw bolt", "10mm", "28", "700", "1500"),
    ("", "Raw bolt", "12mm", "60", "1000", "2000"),
    ("", "Drill spanner", "10mm", "8", "3500", "5000"),
    ("", "Ball gauge", "plastic", "90", "600", "1000"),
    ("", "Ball gauge", "gold", "24", "1000", "1500"),
    ("", "Raw bolt", "6mm", "56", "500", "800"),
    ("Aristo", "Stopper", "Heavy", "78", "2000", "3000"),
    ("Aristo", "Stopper", "Light", "22", "1500", "2500"),
    ("Aristo", "Stopper", "Black 4\"", "23", "2000", "2500"),
    ("Aristo", "Stopper", "Gold", "61", "3000", "4500"),
]
for r in PLUMBING_P6_GEN:
    if len(r) == 4:
        brand, item, spec, qty = r
        PLUMBING.append(row("6", "GENERAL ITEMS", item, brand=brand, category=spec, qty=qty, unit="pcs", notes="Prices not on scan"))
    else:
        brand, item, spec, qty, bp, sp, *u = r
        unit = u[0] if u else "pcs"
        PLUMBING.append(
            row("6", "GENERAL ITEMS", item, brand=brand, category=spec, qty=qty, cost=bp, retail=sp, unit=unit)
        )

PLUMBING_P7 = [
    ("", "Komeo Meco", "", "14", "4000", "5000"),
    ("", "Aristo Gold Stopper Large", "", "10", "5000", "7000"),
    ("", "Komeo Anchor", "", "18", "1500", "2500"),
    ("", "Aristo Stopper Black 3\"", "", "60", "850", "1500"),
    ("Ahadi", "Drawer Lock", "", "33", "1500", "2500"),
    ("Andika", "Drawer Lock", "", "29", "2000", "3000"),
    ("", "Xia boshi Drawer Lock small", "", "18", "2000", "3000"),
    ("", "Xia boshi Drawer Lock Longneck", "", "5", "2600", "3500"),
    ("Aristo", "Cupboard Lock", "", "8", "2500", "3000"),
    ("", "Solex Drawer Lock", "", "11", "4500", "5000"),
    ("", "Door Hinges 2 1/2\"", "", "235", "500", "1000/pair", "pcs @ 500"),
    ("", "Door Hinges 4\"", "", "132", "1500", "3000/pair", "pcs @ 1500"),
    ("", "Door Hinges 3\"", "", "156", "1000", "2000/pair", "pcs @ 1000"),
    ("", "Door Hinges 2\"", "", "307", "500", "1000/pair", "pcs @ 500"),
    ("", "Kobiro kubwa", "", "22", "2500", "5000"),
    ("", "Kobiro ndogo", "", "7", "2000", "4000"),
    ("", "Wire brush yellow", "", "17", "2000", "3500"),
    ("Acacia", "Wire brush red", "", "1", "3500", "5000"),
    ("", "Everray Pawaba - Door Hinges", "", "60", "12000", "17000"),
    ("2.5mm", "Fuxia Door Hinges 4\" * 3\"", "", "2", "3000", "5000", "boxes"),
    ("3 mm", "Fuxia Door Hinges 4\" * 3\"", "", "2", "3000", "5000", "boxes"),
    ("2.5mm", "Fuxia Door Hinges 3\" * 2 1/2\"", "", "3", "3000", "5000", "boxes"),
    ("", "Bed clamp (Heavy duty) 5\"", "", "22", "3000", "5000", "boxes"),
    ("", "Oxford Steel door hinges Silver", "", "43", "3000", "6000", "box"),
    ("", "Spider Door Window", "", "6", "5000", "8000"),
    ("", "Trucker Lock Rubwa", "", "6", "3000", "4000"),
    ("", "Trucker back", "", "", "", "", "Empty row on scan"),
    ("", "plastic Rack", "", "3", "7000", "9000"),
    ("", "Yellow Rack", "", "4", "9500", "15000"),
    ("", "Metal Rack blue Local", "", "3", "5000", "7000"),
]
for brand, item, _x, qty, bp, sp, *extra in PLUMBING_P7:
    if item == "Trucker back":
        PLUMBING.append(row("7", "Locks, hinges & racks", item, brand=brand, notes="Empty row on scan"))
        continue
    unit = "pcs"
    notes = ""
    if extra:
        if extra[-1] in ("boxes", "box"):
            unit = extra[-1]
            notes = " ".join(str(e) for e in extra[:-1]) if len(extra) > 1 else ""
        elif "@" in str(extra[0]) or "pair" in str(sp):
            notes = extra[0]
        else:
            unit = extra[0]
    PLUMBING.append(
        row("7", "Locks, hinges & racks", item, brand=brand, qty=qty, cost=bp, retail=sp, unit=unit, notes=notes)
    )

GENERAL2: list[list] = [
    row("1", "INTERIOR SHELF (11/05/2026)", "Sample Product", code="PRD001", category="General", qty="10", cost="1000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Pipe range 8\"", qty="7", cost="7000", retail="10000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Pipe range 10\"", qty="8", cost="9500", retail="13500", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Pipe range 12\"", qty="1", cost="12500", retail="15000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Pipe range 14\"", qty="5", cost="15000", retail="20000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Pipe range 18\"", qty="3", cost="18000", retail="25000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Adjustable spanner 30mm", qty="7", cost="10500", retail="15000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Adjustable spanner 20mm", qty="6", cost="6500", retail="9000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Sprinkler", qty="2 (was 4)", cost="10000", retail="15000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Panel pins", qty="18", cost="1000", retail="2000", unit="per box"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Spanner 10", cost="2300", retail="3500", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Tiles cleaner (Master) 1L", qty="12", cost="5500", retail="8000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Tiles cleaner 1L", qty="6", cost="5500", retail="7000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Bobo piller Tape", code="DL 008", qty="12", cost="18000", retail="25000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Tangit 1000g", qty="8", cost="8000", retail="15000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Tangit 250g", qty="16", cost="5000", retail="8000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Henkel Tangit 1000ml", qty="3 (was 1)", cost="29000", retail="33000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Tangit 125g", qty="7", cost="1200", retail="3500", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Pipe joint Boss white 200g", qty="5", cost="2000", retail="4000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Water tank Linlee", qty="5", cost="27000", retail="35000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Water tank (Fortec)", qty="1", cost="40000", retail="45000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Angle valve 0.5x0.5\" (Goodone)", qty="6 (was 12)", cost="8000", retail="20000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Angle valve 0.5\" (Wahkit)", qty="1", cost="6000", retail="10000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Angle valve double (Robo)", qty="12", cost="8000", retail="12000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Basket Sink strainer 1 x 1 1/2\"", qty="3+1", cost="15000", retail="25000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "PPR - Conceal tape 3/4\"", qty="4", cost="20000", retail="22000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "PPR - Conceal tape 1/2\"", qty="1", cost="20000", retail="25000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "PPR - Conceal tape 1/2\"", qty="7+3", cost="20000", retail="25000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Thread seal", qty="36", cost="250", retail="500", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Silicone sausage", qty="7", cost="8000", retail="15000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Single Mixer (Best Mix)", category="Kitchen", qty="8", cost="18000", retail="28000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Single Mixer (Blue)", category="Kitchen", qty="3", cost="18000", retail="28000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Gasket Shellac", qty="16", cost="2500", retail="4000", unit="pcs"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Spacer 2.5mm", qty="15", cost="1500", retail="2000", unit="packets"),
    row("1", "INTERIOR SHELF (11/05/2026)", "Spacer 2.0mm", qty="15", cost="1500", retail="2000", unit="packets"),
]

G2_P2_S1 = [
    ("Kitchen single mixer (GFL)", "8", "18000", "25000", "pcs", "Straight crossed out above"),
    ("Shattaf Chuma (Grohe)", "12+1", "10500", "20000", "pcs"),
    ("Shattaf plastic (Grohe)", "4", "6000", "15000", "pcs"),
    ("Pillar cock 1/2' (GFL)", "3", "18000", "22000", "pcs"),
    ("Pillar cock (cobra)", "4", "7000", "12000", "pcs"),
    ("Shower head (Tivoli)", "4", "18000", "25000", "pcs"),
    ("Shower head L108", "3+1", "12000", "17000", "pcs"),
    ("Shower head L175", "6+1", "8000", "12000", "pcs"),
    ("Shower head ndogo (Nobranded)", "5", "2500", "5000", "pcs"),
    ("Kizoleo", "6", "1000", "2000", "pcs"),
    ("Andika Cabinet Hinges", "189", "908", "2000", "pairs"),
    ("Andika cabinet Hinges hydraulic", "54", "1640", "3000", "pairs"),
    ("Cabinet handles (Ndogo)", "6+18+1", "800", "1500", "pcs"),
    ("Cable ties Ndogo", "4", "5000", "8000", "packs"),
    ("Cable ties Kubwa", "2", "6000", "12000", "packs"),
    ("Zibulio", "6", "2500", "5000", "pcs"),
    ("Cabinet Handles (Kubwa)", "15", "900", "2500", "pcs"),
    ("Cabinet Handles (Kati)", "28", "850", "2000", "pcs"),
    ("Cabinet Handles (Standard)", "14+6", "800", "1500", "pcs"),
    ("Soap dish", "1+1", "5000", "10000", "pcs"),
    ("Toilet paper handle", "1", "5000", "10000", "pcs"),
]
for item, qty, c, r, u, *n in G2_P2_S1:
    GENERAL2.append(row("2", "Kitchen & cabinet (top)", item, qty=qty, cost=c, retail=r, unit=u, notes=n[0] if n else ""))

G2_P2_S2 = [
    ("Mallet Hammer", "5", "7000", "7000", "pcs", "Written over 5000"),
    ("Kibao Decoration", "14", "2500", "5000", "pcs"),
    ("Chesal patasi (Acassia) 19mm", "4", "7500", "9500", "pcs"),
    ("Chesal patasi (Acassia) 12mm", "4", "6500", "7500", "pcs"),
]
for item, qty, c, r, u, *n in G2_P2_S2:
    GENERAL2.append(row("2", "General (11/05/20)", item, qty=qty, cost=c, retail=r, unit=u, notes=n[0] if n else ""))

G2_P2_X = [
    ("Fevicol 1kg pur 250g", "19", "800", "12000"),
    ("Fevicol 1kg pur 125g", "8", "500", "6500"),
    ("Fevicol 1kg pur 500g", "6", "1700", "19000"),
    ("Basco paint Light Oak Varnish 1L", "11", "4500", "6000"),
    ("Basco paint Varnish clear 1L", "13", "4500", "6000"),
    ("Mamba Highgloss thinner 5L", "6", "30000", "35000"),
    ("Mamba Solvent 5L", "7", "25000", "30000"),
    ("Mamba Standard thinner 5L", "3", "22500", "25000"),
]
for item, qty, c, r in G2_P2_X:
    GENERAL2.append(row("2", "General (11/05/20) — CROSSED OUT", item, qty=qty, cost=c, retail=r, unit="pcs", notes="Large X on scan — verify before use"))

G2_P3 = [
    ("Lamp Holder (Eurotrix)", "30+1", "", "", "pcs"),
    ("Ceiling rose with lamp holder", "Hippo", "5pc+2", "4000", "pcs", "Brand in qty column"),
    ("Extension Tronic 1.25mm²", "5", "", "25000", "pcs"),
    ("Africab Lampholder", "4", "", "", "pcs"),
    ("Single switch socket (Torch)", "6+1", "", "10000", "pcs"),
    ("USB Single switch socket (Eurotrix)", "", "", "", ""),
    ("1g 2w Switch", "4", "", "2500", "pcs"),
    ("DP switch 45A Torch", "8+1", "", "7000", "pcs"),
    ("DP switch 20A Torch", "3", "", "6000", "pcs"),
    ("DP switch 20A Maxipower", "8", "", "6000", "pcs"),
    ("2g 2w Switch (Torch)", "7", "", "", ""),
    ("DP socket + USB (Torch)", "6", "", "15000", "pcs"),
    ("Africab Ceiling rose with holder", "5", "", "5000", "pcs"),
    ("Africab 2g 1w switch", "12+8+5", "1500", "3500", "pcs"),
    ("4way switch (Torch)", "2", "2000", "4000", "pcs"),
    ("Africab 1g 2w switch", "3", "", "3500", "pcs"),
    ("Africab 1g 1w switch", "21", "", "", "pcs"),
    ("Africab 2g 2w switch", "20", "", "3500", "pcs"),
    ("Africab switch socket", "7", "", "3500", "pcs", "Retail crossed out"),
    ("Africab switch socket double", "1", "", "6000", "pcs"),
    ("Africab 3g 2w switch", "2+10+2", "2000", "3500", "pcs"),
    ("4g 2w switch Tronic", "6", "", "5000", "pcs"),
    ("DP switch Tronic 20A", "3", "", "8000", "pcs"),
    ("Junction box", "20", "", "2000", "pcs"),
    ("Multiplug (Eurotrix)", "21", "", "", "pcs"),
    ("Lampholder (Alivons)", "20", "", "", "pcs"),
    ("DP switch socket 13A", "5", "", "6000", "pcs"),
    ("Legrand Double socket", "5", "", "45000", "pcs"),
    ("Switch cooker (Tronic)", "8", "", "30000", "pcs"),
    ("2g 2w switch (Maxipower)", "6+1", "", "3500", "pcs"),
]
for item, *rest in G2_P3:
    qty = rest[0] if rest else ""
    cost = rest[1] if len(rest) > 1 else ""
    retail = rest[2] if len(rest) > 2 else ""
    unit = rest[3] if len(rest) > 3 else "pcs"
    notes = rest[4] if len(rest) > 4 else ""
    GENERAL2.append(row("3", "ELECTRICAL SHELF", item, qty=qty, cost=cost, retail=retail, unit=unit, notes=notes))

for tup in [
    ("Africab Top plug", "32", "", "3000"),
    ("Africab cover single", "16", "", "2000"),
    ("Cover single switch", "192", "", "200", "48×4=192 margin note"),
    ("2way switch", "4", "", "3000"),
    ("Single switch", "5", "", "3000"),
    ("Bulb 24W Vallight", "2", "", "10000"),
    ("Bulb 15W Vallight", "10+1", "", "9000"),
    ("Bulb 25W Vallight", "5", "", "12000"),
    ("Bulb 35W Vallight", "2", "", "15000"),
    ("Alyons Bulb 28W", "5", "", "10000"),
    ("Alyons Bulb 38W", "1", "", "12000"),
    ("Bulb 12W Vallight", "1+2+1", "", "7000"),
    ("Bulb 12W Oliver charging", "1", "", "15000"),
    ("Bulb 12W Vellmax", "6", "", "15000"),
    ("Bulb 9W", "2", "", "9500"),
    ("Bulb 7W Vallight", "41+12", "", "4000"),
    ("Bulb 5W Vallight", "25", "", "3000"),
    ("Bulb 5W", "49", "", "3000"),
    ("Bulb 3W", "54", "", "2500"),
    ("Africab photocell", "6", "", "25000", "Was 15000, crossed out"),
    ("Distribution Box Alyons", "2", "", "", "Asterisk — no price on scan"),
    ("IP4 Way Mainswitch", "1 (crossed)", "", "55000", "Qty/45k crossed"),
    ("Main Switch IP6WAY", "1", "", "65000", "55k crossed"),
    ("Africab Distribution Box", "1", "", "", "Asterisk — no price"),
    ("Distribution Box (Torch)", "2", "", "", "Asterisk — no price"),
    ("Pyramid socket (tronic)", "2", "", "", "Asterisk — no price"),
    ("Circuit Breaker (Africab)", "5+1", "", "45000"),
    ("Circuit Breaker", "1", "", "35000"),
    ("Plascon Emulsion 20L", "15", "", "", "No price on scan"),
    ("Welding Rod E6013 12.5kg", "25", "", "", "Entire row crossed out"),
    ("Miniature 32A", "2", "", "8000"),
    ("Insulation tape", "17", "800", "1000", "800 crossed, 1000 final"),
]:
    item, qty, cost, retail, *n = tup
    GENERAL2.append(row("4", "Electrical (cont.)", item, qty=qty, cost=cost, retail=retail, unit="pcs", notes=n[0] if n else ""))

G2_P5 = [
    ("U-Clip 18mm", "4", "", "", "Box"),
    ("U-clip 22mm", "14", "", "", "Box"),
    ("U-clip 9mm", "12", "1500", "3000", "pcs box"),
    ("Waterproof box IP65", "5", "", "7000", "Pc"),
    ("Waterproof box IP55", "9", "", "5000", "Pc"),
    ("Blue Wire Single Core 1.5mm (Africal)", "2", "65000", "80000", "rolls", "100m"),
    ("L/Blue Wire Single core 1.5mm (Master cable)", "2", "55000", "75000", "rolls", "100m"),
    ("Brown Wire single core 1.5mm (Master)", "2", "55000", "75000", "rolls", "100m"),
    ("Green Wire single core 2.5mm (Africal)", "1", "65000", "80000", "rolls", "100m"),
    ("Wire double 1.5mm (Afuzab)", "1", "", "195000", "roll", "100meter"),
    ("Wire double 1.5mm (Master Cable)", "1", "70000", "95000", "roll"),
    ("Wire double 2.5mm (Master Cable)", "", "", "", ""),
    ("Solar Cable", "120", "", "1000", "meters"),
    ("Euro 2.5mm wire twin", "100", "", "", "meters"),
    ("Euro 1.5mm wire twin", "70", "", "", "meters"),
    ("Switch box Nzito Single", "60+5", "", "", "pcs"),
    ("Switch box Nyepeo Single", "10+15", "", "", ""),
    ("Switch box Nzito Double", "65", "", "", "pc"),
    ("Switch box Nyepeo Double", "31", "", "", "pc"),
    ("Bulb Watts 200", "17", "", "", "pc"),
    ("Bulb watts 100", "7", "", "", "pc"),
    ("Square box Double white", "9", "", "2000", "pc"),
    ("Square box Double Black", "15", "", "2500", "pc"),
    ("Square box Single", "27", "", "1000", "pc"),
    ("Wire black flexible 0.5", "100", "", "1000", "meters"),
]
for item, qty, cost, retail, unit, *n in G2_P5:
    GENERAL2.append(row("5", "Electrical wire & boxes (12/05/2026)", item, qty=qty, cost=cost, retail=retail, unit=unit, notes=n[0] if n else ""))

G2_P6 = [
    ("Bekeshi (Mbao)", "", "8", "11", "", "Crossed out"),
    ("Bekeshi (Mbao)", "", "13", "4", "6", "Crossed out"),
    ("Bekeshi (Mbao) Nlogo", "", "", "", "", "Partially crossed out"),
    ("Choo.", "", "17", "", ""),
    ("Conduit Elbow 3/4\"", "54", "500", "500", "pc"),
    ("Conduit Elbow 1\"", "", "1500", "1500", "pc"),
    ("Conduit Elbow 1\" Supi", "", "1000", "1000", "pc"),
    ("Conduit Elbow 3/4\" Supi", "119", "500", "500, 1000", "pe"),
    ("Connector 1\"", "", "500", "500", "pe"),
    ("Connector 3/4\"", "", "4000", "10000", "pe"),
    ("Junction Box 3/4\"", "56", "", "1000", "pc"),
    ("Junction Box 1\"", "50", "", "1000", "pc"),
    ("Screw & fisher plug 4\"", "119.", "", "1000", "pe"),
    ("Screw & fisher plug 3\"", "54+80", "", "500", "pe"),
    ("Screw & fisher plug 3\" Nyamba", "72+100", "", "300", "pe"),
    ("Screw fisher plug 2 1/2\" Nyamba", "81", "", "", ""),
    ("Cabinet Handle (Stainless steel) Silver medium + Black", "Bx4", "1500 per pc", "", "Bx 4", "Steel 1 Bx, Black 3 Bx — 80pcs"),
    ("Black finish (small)", "Bx2", "1500 per pc", "", "Bx 2", "40 pcs @20pc"),
    ("Stain finish (small)", "Bx2", "1500 per pc", "", "Bx 2", "40pc"),
    ("Black finish (large)", "Bx1", "2500 per pc", "", "Bx 1", "20pc @20pc"),
]
for item, qty, cost, retail, unit, *n in G2_P6:
    GENERAL2.append(row("6", "Conduit & handles (market rate)", item, qty=qty, cost=cost, retail=retail, unit=unit, notes="; ".join(n) if n else ""))


def style_sheet(ws, title: str, rows: list[list]) -> None:
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="EA580C")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.append([])
    ws.append(HEADERS)
    header_row = 3
    fill = PatternFill("solid", fgColor="1E293B")
    font = Font(bold=True, color="F1F5F9")
    thin = Side(style="thin", color="334155")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(header_row, col)
        cell.fill = fill
        cell.font = font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    alt = PatternFill("solid", fgColor="F8FAFC")
    for r in rows:
        ws.append(r)
        row_idx = ws.max_row
        for col in range(1, len(HEADERS) + 1):
            c = ws.cell(row_idx, col)
            c.border = border
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if row_idx % 2 == 0:
                c.fill = alt

    widths = [6, 28, 10, 12, 36, 18, 12, 16, 16, 10, 24]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"


# --- Supafundi import format (matches lib/excel/inventory-columns.ts) ---
IMPORT_HEADERS = [
    "Page",
    "Code",
    "Name",
    "Category",
    "Quantity",
    "Cost",
    "Retail Price",
    "Unit",
]

IMPORT_PLUMBING = Path(
    r"c:\Users\HP\OneDrive\Desktop\Plumbing_and_General_Import.xlsx"
)
IMPORT_GENERAL2 = Path(
    r"c:\Users\HP\OneDrive\Desktop\General_2_Import.xlsx"
)

SKIP_NAME_PARTS = ("sample product", "example nail")
SKIP_CODES = frozenset({"PRD001", "PRD-001"})


def sum_quantity(value: str) -> int | float:
    """Match Supafundi parseQuantityCell: sum tallies like 7+1, 12+8+5."""
    s = str(value or "").strip()
    if not s:
        return 0
    if "(" in s:
        s = s.split("(")[0].strip()
    s = re.sub(r"\bpcs?\.?\b", "", s, flags=re.I).strip()
    s = s.replace(",", "")
    if re.fullmatch(r"\d+(\.\d+)?", s):
        n = float(s)
        return int(n) if n == int(n) else n
    if "+" in s:
        total = 0.0
        for part in s.split("+"):
            part = re.sub(r"[^\d.]", "", part.strip())
            if not part:
                continue
            n = float(part)
            if n >= 0:
                total += n
        if total > 0:
            return int(total) if total == int(total) else total
    digits = re.sub(r"[^\d.]", "", s)
    if digits:
        n = float(digits)
        return int(n) if n == int(n) else n
    return 0


def parse_price(value: str) -> int | float | "":
    s = str(value or "").strip()
    if not s:
        return ""
    if "/" in s:
        s = s.split("/")[0].strip()
    s = s.replace(",", "")
    m = re.search(r"\d+(?:\.\d+)?", s)
    if not m:
        return ""
    n = float(m.group())
    return int(n) if n == int(n) else n


def infer_category(section: str, category: str, name: str) -> str:
    cat = category.strip()
    if cat:
        low = cat.lower()
        if low == "plumbing":
            return "Plumbing"
        if low == "kitchen":
            return "Kitchen"
        if low == "general":
            return "General"
        return cat
    sec = section.lower()
    name_low = name.lower()
    if "electrical" in sec or "wire" in sec and "pipe" not in name_low:
        return "Electrical"
    if "kitchen" in sec:
        return "Kitchen"
    if any(
        k in sec
        for k in ("plumbing", "pvc", "ppr", "ips", "gutter", "fitting", "front desk")
    ):
        return "Plumbing"
    if any(k in name_low for k in ("ppr ", "pvc ", "ips ", "gs ", "pex ", "elbow", "tee ")):
        return "Plumbing"
    return "General"


def normalize_unit(unit: str) -> str:
    u = str(unit or "").strip().lower()
    if not u or u in ("per", "pe", "pc"):
        return "pcs"
    if u.startswith("pc"):
        return "pcs"
    if u in ("box", "boxes"):
        return "box"
    if u in ("pack", "packs"):
        return "pack"
    if u in ("pair", "pairs"):
        return "pair"
    if u in ("roll", "rolls", "roller"):
        return "roll"
    if u in ("meter", "meters"):
        return "meter"
    if "packet" in u:
        return "packet"
    return str(unit or "pcs").strip() or "pcs"


def raw_to_import_row(raw: list) -> list | None:
    page, section, code, brand, item, category, qty, cost, retail, unit, *_notes = (
        raw + [""] * 11
    )[:11]
    name = str(item).strip()
    if not name:
        return None
    name_low = name.lower()
    if any(skip in name_low for skip in SKIP_NAME_PARTS):
        return None
    code_s = str(code).strip()
    if code_s.upper().replace(" ", "") in SKIP_CODES:
        return None
    if brand:
        brand_s = str(brand).strip()
        if brand_s and brand_s.lower() not in name_low:
            name = f"{brand_s} {name}"
    cat = infer_category(str(section), str(category), name)
    quantity = sum_quantity(str(qty))
    cost_p = parse_price(str(cost))
    retail_p = parse_price(str(retail))
    unit_out = normalize_unit(str(unit))
    page_out = int(page) if str(page).isdigit() else page
    return [
        page_out,
        code_s,
        name,
        cat,
        quantity,
        cost_p if cost_p != "" else "",
        retail_p if retail_p != "" else "",
        unit_out,
    ]


def write_supafundi_import(path: Path, raw_rows: list[list]) -> int:
    import_rows = []
    for raw in raw_rows:
        row_out = raw_to_import_row(raw)
        if row_out:
            import_rows.append(row_out)
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock List"
    ws.append(IMPORT_HEADERS)
    for r in import_rows:
        ws.append(r)
    widths = [6, 12, 42, 14, 10, 12, 12, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return len(import_rows)


def main() -> None:
    n1 = write_supafundi_import(IMPORT_PLUMBING, PLUMBING)
    n2 = write_supafundi_import(IMPORT_GENERAL2, GENERAL2)
    print(f"Saved: {IMPORT_PLUMBING} ({n1} rows)")
    print(f"Saved: {IMPORT_GENERAL2} ({n2} rows)")


if __name__ == "__main__":
    main()
